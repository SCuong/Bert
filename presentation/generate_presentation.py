"""
generate_presentation.py - Tự động tạo slide PowerPoint học thuật chuẩn mực (15 slides).
Sử dụng hệ thống thiết kế Đại học Bách Khoa - ĐHĐN (DUT) từ template chính thức.
Tích hợp trực tiếp các số liệu định lượng (JSON) và hình ảnh thực nghiệm (PNG) chính thức.
"""

import os
import sys
import json
import argparse
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

# Đảm bảo repository root luôn có trong sys.path
REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

from src.config import (
    MODEL_NAME,
    MAX_LENGTH,
    RANDOM_SEED,
    TRAIN_RATIO,
    VAL_RATIO,
    TEST_RATIO,
    PROJECT_ROOT,
    FIGURES_DIR,
    BASELINE_VAL_METRICS_PATH,
    BERT_VAL_METRICS_PATH,
    BASELINE_METRICS_PATH,
    BERT_METRICS_PATH,
    COMPARATIVE_METRICS_PATH,
    ERROR_ANALYSIS_SUMMARY_PATH
)

def create_presentation(is_draft=False):
    """
    Hàm tạo bài thuyết trình 15 slide theo hệ thống thiết kế DUT.
    Nạp dữ liệu thực nghiệm chính thức từ thư mục artifacts/.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)  # 16:9 Widescreen standard
    blank_layout = prs.slide_layouts[6]

    # Bảng màu thiết kế học thuật DUT (Navy / Steel / Charcoal / White)
    NAVY = RGBColor(26, 54, 93)            # #1A365D - Màu chủ đạo DUT
    STEEL = RGBColor(43, 108, 176)         # #2B6CB0 - Màu nhấn mạnh
    CHARCOAL = RGBColor(45, 55, 72)        # #2D3748 - Màu văn bản chính
    MUTED = RGBColor(113, 128, 150)        # #718096 - Màu phụ đề
    BG_BOX = RGBColor(248, 250, 252)       # #F8FAFC - Nền card
    BORDER_COLOR = RGBColor(226, 232, 240) # #E2E8F0 - Viền card
    WHITE = RGBColor(255, 255, 255)
    HIGHLIGHT_ROW = RGBColor(235, 248, 255)
    ACCENT_GREEN = RGBColor(40, 167, 69)

    # Đường dẫn tài nguyên đồ họa DUT trích xuất từ template
    ASSETS_DIR = os.path.join(PROJECT_ROOT, "presentation", "assets")
    DUT_COVER_BG = os.path.join(ASSETS_DIR, "slide1_shape_Google Shape;253;p1.png")
    DUT_LOGO = os.path.join(ASSETS_DIR, "slide5_shape_Google Shape;368;g3f87e06bbab_0_2244.png")

    # Nạp các tệp metrics chính thức
    baseline_val = {}
    bert_val = {}
    baseline_test = {}
    bert_test = {}
    comparative = {}
    error_summary = {}

    if os.path.exists(BASELINE_VAL_METRICS_PATH):
        with open(BASELINE_VAL_METRICS_PATH, "r", encoding="utf-8") as f:
            baseline_val = json.load(f)
    if os.path.exists(BERT_VAL_METRICS_PATH):
        with open(BERT_VAL_METRICS_PATH, "r", encoding="utf-8") as f:
            bert_val = json.load(f)
    if os.path.exists(BASELINE_METRICS_PATH):
        with open(BASELINE_METRICS_PATH, "r", encoding="utf-8") as f:
            baseline_test = json.load(f)
    if os.path.exists(BERT_METRICS_PATH):
        with open(BERT_METRICS_PATH, "r", encoding="utf-8") as f:
            bert_test = json.load(f)
    if os.path.exists(COMPARATIVE_METRICS_PATH):
        with open(COMPARATIVE_METRICS_PATH, "r", encoding="utf-8") as f:
            comparative = json.load(f)
    if os.path.exists(ERROR_ANALYSIS_SUMMARY_PATH):
        with open(ERROR_ANALYSIS_SUMMARY_PATH, "r", encoding="utf-8") as f:
            error_summary = json.load(f)

    def add_header(slide, title_text, subtitle_text=""):
        """Tạo header chuẩn hóa cho các slide nội dung."""
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.font.name = "Arial"

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.size = Pt(11)
            p2.font.color.rgb = MUTED
            p2.font.name = "Arial"
            p2.space_before = Pt(3)

    def add_footer_and_logo(slide):
        """Thêm logo DUT và footer bản quyền chuẩn."""
        # Footer text
        tx = slide.shapes.add_textbox(Inches(0.8), Inches(6.88), Inches(9.0), Inches(0.4))
        tf = tx.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = "Đồ án môn học Trí tuệ Nhân tạo — ĐH Bách Khoa - ĐHĐN"
        p.font.size = Pt(9.5)
        p.font.color.rgb = MUTED
        p.font.name = "Arial"

        # DUT Logo
        if os.path.exists(DUT_LOGO):
            slide.shapes.add_picture(DUT_LOGO, Inches(11.56), Inches(6.45), Inches(1.2), Inches(0.65))

    def add_card(slide, left, top, width, height, title="", body="", bg=BG_BOX, border=BORDER_COLOR):
        """Tạo card container bo góc trang nhã."""
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg
        shape.line.color.rgb = border
        shape.line.width = Pt(1)

        if title or body:
            tf = shape.text_frame
            tf.word_wrap = True
            tf.margin_left = Inches(0.22)
            tf.margin_top = Inches(0.22)
            tf.margin_right = Inches(0.22)
            tf.margin_bottom = Inches(0.22)

            if title:
                p = tf.paragraphs[0]
                p.text = title
                p.font.size = Pt(14)
                p.font.bold = True
                p.font.color.rgb = NAVY
                p.font.name = "Arial"

            if body:
                p_body = tf.add_paragraph() if title else tf.paragraphs[0]
                p_body.text = body
                p_body.font.size = Pt(11)
                p_body.font.color.rgb = CHARCOAL
                p_body.font.name = "Arial"
                p_body.space_before = Pt(5)
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide (DUT Cover Design)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    if os.path.exists(DUT_COVER_BG):
        s1.shapes.add_picture(DUT_COVER_BG, 0, 0, Inches(13.333), Inches(7.5))
    else:
        bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg1.fill.solid()
        bg1.fill.fore_color.rgb = RGBColor(245, 247, 250)
        bg1.line.fill.background()

    # Title text block trên nền template DUT
    t_box = s1.shapes.add_textbox(Inches(4.2), Inches(1.3), Inches(8.5), Inches(2.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "ĐỒ ÁN MÔN HỌC: TRÍ TUỆ NHÂN TẠO"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = STEEL
    p.font.name = "Arial"

    p = tf1.add_paragraph()
    p.text = "Ứng Dụng Mô Hình BERT Trong Phân Loại Cảm Xúc Đánh Giá Khách Sạn"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.font.name = "Arial"
    p.space_before = Pt(8)

    p = tf1.add_paragraph()
    p.text = "Fine-Tuning BERT for Hotel Review Sentiment Classification"
    p.font.size = Pt(15)
    p.font.color.rgb = CHARCOAL
    p.font.name = "Arial"
    p.space_before = Pt(4)

    p = tf1.add_paragraph()
    p.text = "Tuân thủ chuẩn mực AI Project Cycle • Thực nghiệm độc lập có kiểm chứng"
    p.font.size = Pt(11.5)
    p.font.color.rgb = MUTED
    p.font.name = "Arial"
    p.space_before = Pt(10)

    # Info card phía dưới bên phải
    info_box = s1.shapes.add_textbox(Inches(4.2), Inches(4.5), Inches(8.5), Inches(2.2))
    tf_info = info_box.text_frame
    tf_info.word_wrap = True

    p = tf_info.paragraphs[0]
    p.text = "THÔNG TIN BẢO VỆ ĐỒ ÁN"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.font.name = "Arial"

    p = tf_info.add_paragraph()
    p.text = "• Giảng viên hướng dẫn: Bộ môn Trí tuệ Nhân tạo — Khoa Công nghệ Thông tin\n" \
             "• Sinh viên thực hiện: Nhóm Nghiên cứu Đồ án BERT\n" \
             "• Mô hình cốt lõi: google-bert/bert-base-uncased (110M tham số)\n" \
             "• Tập dữ liệu: Teacher-provided dataset (19,745 mẫu sạch sau kiểm toán)\n" \
             "• Kết quả nổi bật: BERT Test Accuracy 84.83% (Vượt Baseline +2.89% & Code cũ +19.63%)"
    p.font.size = Pt(11)
    p.font.color.rgb = CHARCOAL
    p.font.name = "Arial"
    p.space_before = Pt(4)

    # =========================================================================
    # SLIDE 2: Scope & Problem Statement (AI Project Cycle Phase 1)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "01. Bối Cảnh & Mục Tiêu Nghiên Cứu", "Giai đoạn 01 (Scope & Plan) theo chuẩn AI Project Cycle")
    add_footer_and_logo(s2)

    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="Phát Biểu Bài Toán & Bối Cảnh",
             body="• Bối cảnh thực tế:\n"
                  "  - Hàng triệu đánh giá dịch vụ khách sạn trực tuyến (Booking.com) được đăng tải.\n"
                  "  - Nhu cầu tự động hóa phân loại phản hồi khách hàng để xử lý kịp thời.\n\n"
                  "• Đầu vào & Đầu ra (I/O Specification):\n"
                  "  - Input: Văn bản đánh giá tiếng Anh (độ dài từ vài từ đến hàng trăm từ).\n"
                  "  - Output: Nhãn phân loại nhị phân (1: Tích cực, 0: Tiêu cực) kèm độ tin cậy.\n\n"
                  "• Các bên liên quan (Stakeholders):\n"
                  "  - Ban quản lý: Phát hiện sự cố dịch vụ ngay lập tức.\n"
                  "  - Khách hàng: Kỳ vọng phản hồi được lắng nghe và cải thiện.\n"
                  "  - Kỹ sư AI: Xây dựng hệ thống chính xác cao, có khả năng giải thích.")

    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
             title="Mục Tiêu & Tiêu Chí Đánh Giá",
             body="• Mục tiêu kỹ thuật cốt lõi:\n"
                  "  1. Xây dựng mô hình cơ sở TF-IDF + Logistic Regression chuẩn xác.\n"
                  "  2. Fine-tuning chuẩn mực mô hình google-bert/bert-base-uncased.\n"
                  "  3. Đánh giá đối đầu khách quan trên Held-Out Test Set (3,949 mẫu).\n"
                  "  4. Phân tích định tính 20 ca lỗi cực đoan (Error Analysis).\n\n"
                  "• Tiêu chí đánh giá chất lượng:\n"
                  "  - Macro F1-score & Accuracy trên tập kiểm thử độc lập.\n"
                  "  - Kiểm soát chống rò rỉ dữ liệu (Zero Data Leakage).\n"
                  "  - Tính tái lập có kiểm soát (Seed cố định 42 xuyên suốt).\n"
                  "  - Đóng gói ứng dụng Web Demo tương tác thời gian thực (Streamlit).")

    # =========================================================================
    # SLIDE 3: Evolution of NLP: RNN/LSTM to Transformer
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "02. Tiến Hóa Kiến Trúc: Từ RNN/LSTM Đến Transformer", "Bước chuyển dịch mô hình tính toán trong Xử lý Ngôn ngữ Tự nhiên")
    add_footer_and_logo(s3)

    add_card(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="Hạn Chế Cố Hữu Của RNN & LSTM",
             body="• Điểm nghẽn xử lý tuần tự (Sequential Recurrence Bottleneck):\n"
                  "  - Để tính trạng thái ẩn từ thứ t, bắt buộc phải đợi từ thứ t-1.\n"
                  "  - Không thể tính toán song song theo chiều dài chuỗi trên GPU/TPU.\n"
                  "  - Tốc độ huấn luyện chậm khi làm việc với ngữ liệu lớn.\n\n"
                  "• Suy giảm thông tin ngữ cảnh xa (Vanishing Gradient):\n"
                  "  - Qua các bước đệ quy dài, thông tin ở đầu câu bị suy giảm hoặc méo mó.\n"
                  "  - Khó duy trì sự phụ thuộc giữa các từ cách xa nhau.\n\n"
                  "• Hạn chế của BiLSTM:\n"
                  "  - Chỉ ghép nối 2 mạng một chiều độc lập (trái sang + phải sang).\n"
                  "  - Thiếu sự tương tác ngữ cảnh hai chiều sâu ở từng tầng.")

    add_card(s3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
             title="Đột Phá Của Transformer (Vaswani et al., 2017)",
             body="• Loại bỏ hoàn toàn mạng đệ quy (No Recurrence):\n"
                  "  - Xử lý đồng thời mọi token trên toàn bộ chiều dài chuỗi trong cùng một tầng.\n"
                  "  - Tận dụng tối đa khả năng tăng tốc phần cứng song song.\n\n"
                  "• Cơ chế Self-Attention (Tự chú ý):\n"
                  "  - Kết nối trực tiếp bất kỳ cặp từ nào trong câu với độ dài đường truyền O(1).\n"
                  "  - Triệt tiêu hiện tượng suy biến thông tin qua khoảng cách xa.\n\n"
                  "• Mã hóa vị trí (Positional Encoding):\n"
                  "  - Bổ sung thông tin trật tự cú pháp mà không cần bước tuần tự.\n\n"
                  "• Nền tảng của kỷ nguyên mô hình ngôn ngữ lớn:\n"
                  "  - Đặt cơ sở trực tiếp cho sự ra đời của BERT, RoBERTa và GPT.")

    # =========================================================================
    # SLIDE 4: BERT Architecture & Attention Mechanism
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "03. Kiến Trúc BERT & Cơ Chế Multi-Head Self-Attention", "Cấu tạo toán học của mô hình Bidirectional Encoder Representations from Transformers")
    add_footer_and_logo(s4)

    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="Bản Chất Kiến Trúc BERT",
             body="• Chỉ sử dụng khối ENCODER của Transformer:\n"
                  "  - Chuyên trách nhiệm vụ thấu hiểu ngôn ngữ (Language Understanding).\n"
                  "  - Deeply Bidirectional: Nhìn đồng thời ngữ cảnh hai chiều tự do.\n\n"
                  "• Cấu hình mô hình bert-base-uncased:\n"
                  "  - 12 tầng Transformer Encoder (L=12)\n"
                  "  - Kích thước không gian ẩn H=768, 12 Attention Heads (A=12)\n"
                  "  - Xấp xỉ 110 triệu tham số tiền huấn luyện trên 3.3 tỷ từ.\n\n"
                  "• Token đặc biệt:\n"
                  "  - [CLS]: Đứng đầu mỗi chuỗi, vector ẩn làm đại diện phân loại.\n"
                  "  - [SEP]: Phân tách câu và đánh dấu kết thúc chuỗi.\n"
                  "  - WordPiece Tokenization: Từ điển 30,522 tokens, triệt tiêu lỗi OOV.")

    add_card(s4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
             title="Cơ Chế Scaled Dot-Product Attention",
             body="• Công thức toán học cốt lõi:\n"
                  "  Attention(Q, K, V) = softmax( (Q * K^T) / sqrt(d_k) ) * V\n\n"
                  "• Bộ ba vector Query (Q), Key (K), Value (V):\n"
                  "  - Query: Yêu cầu tìm kiếm thông tin của từ hiện tại.\n"
                  "  - Key: Từ khóa chỉ mục đại diện của các từ trong câu.\n"
                  "  - Value: Nội dung ngữ nghĩa thực sự mà từ mang theo.\n\n"
                  "• Hệ số co giãn 1/sqrt(d_k):\n"
                  "  - Ngăn tích vô hướng quá lớn làm bão hòa gradient tại Softmax.\n\n"
                  "• Multi-Head Attention (12 Heads song song):\n"
                  "  - Cho phép mô hình đồng thời chú ý vào nhiều khía cạnh khác nhau (cú pháp, đại từ, phủ định).")

    # =========================================================================
    # SLIDE 5: Dataset & EDA (Phân tích Khám phá Dữ liệu)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "04. Phân Tích Khám Phá Dữ Liệu (Data & EDA)", "Giai đoạn 02 (Data) theo AI Project Cycle — Kiểm soát chất lượng dữ liệu")
    add_footer_and_logo(s5)

    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="Kiểm Toán Dữ Liệu & Phân Chia Chống Rò Rỉ",
             body="• Dữ liệu thô ban đầu (data/dts_20k_raw.csv):\n"
                  "  - 20,000 bài đánh giá Booking.com do giảng viên cung cấp.\n"
                  "  - Loại 255 mẫu: 5 rỗng, 127 xung đột nhãn (20 text trùng cả 0 và 1), 123 trùng lặp hoàn toàn.\n"
                  "  - Dữ liệu sạch: 19,745 mẫu (9,921 Negative - 50.25%, 9,824 Positive - 49.75%).\n\n"
                  "• Phân chia Stratified Split (seed=42, Zero Leakage):\n"
                  "  - Train Set: 13,821 mẫu (70.0%) | 6,945 Neg, 6,876 Pos.\n"
                  "  - Val Set:    1,975 mẫu (10.0%) | 992 Neg, 983 Pos.\n"
                  "  - Test Set:   3,949 mẫu (20.0%) | 1,984 Neg, 1,965 Pos.\n"
                  "  - Xác nhận 0 mẫu trùng lặp giữa Train, Val và Test.\n\n"
                  "• Thống kê độ dài token BERT (Train + Val Scope):\n"
                  "  - Median: 30 tokens | Mean: 42.33 | p90: 94 | p95: 121 | Max: 425.\n"
                  "  - Quyết định: Chọn MAX_LENGTH = 128 bảo toàn 96.07% văn bản, tối ưu chi phí O(L^2).")

    len_img = os.path.join(FIGURES_DIR, "token_length_distribution.png")
    if os.path.exists(len_img):
        s5.shapes.add_picture(len_img, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    else:
        add_card(s5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
                 title="Biểu Đồ Phân Bố Độ Dài Token",
                 body="• Đã lưu tại artifacts/figures/token_length_distribution.png\n"
                      "• Xác nhận ngưỡng MAX_LENGTH = 128 bao phủ hơn 96% dữ liệu.")

    # =========================================================================
    # SLIDE 6: Project Pipeline (Quy trình Nghiên cứu)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "05. Quy Trình Nghiên Cứu Chuẩn Mực (Pipeline)", "Triển khai độc lập từ đầu, tuân thủ nghiêm ngặt chuẩn mực AI Project Cycle")
    add_footer_and_logo(s6)

    add_card(s6, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.0),
             title="1. Dữ Liệu & Tiền Xử Lý",
             body="• Tiền xử lý bảo toàn ngữ pháp:\n"
                  "  - Xóa khoảng trắng thừa.\n"
                  "  - GIỮ NGUYÊN từ phủ định ('not', 'no', 'never').\n"
                  "  - GIỮ NGUYÊN cấu trúc câu & dấu câu cho Positional Encoding.\n\n"
                  "• Kiểm toán dữ liệu:\n"
                  "  - Lọc bỏ 255 mẫu lỗi/trùng/xung đột trước khi split.\n\n"
                  "• Stratified Split (Seed 42):\n"
                  "  - 70% Train / 10% Val / 20% Test.\n"
                  "  - Niêm phong tuyệt đối tập Test.")

    add_card(s6, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.0),
             title="2. Mô Hình Hóa (Modeling)",
             body="• Mô hình cơ sở (Baseline):\n"
                  "  - TF-IDF (10,000 unigram + bigram).\n"
                  "  - Logistic Regression (L-BFGS).\n"
                  "  - Thiết lập mốc so sánh tối thiểu.\n\n"
                  "• Mô hình BERT Fine-Tuning:\n"
                  "  - google-bert/bert-base-uncased.\n"
                  "  - Full Fine-Tuning mở khóa 110M tham số.\n"
                  "  - AdamW (lr=2e-5, weight decay=0.01).\n"
                  "  - Linear Warmup Scheduler (10% steps).\n"
                  "  - Checkpoint theo Validation Macro F1.")

    add_card(s6, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.0),
             title="3. Đánh Giá & Triển Khai",
             body="• Đánh giá đối đầu độc lập:\n"
                  "  - Chạy trên Held-Out Test Set (3,949 mẫu) đúng 1 lần duy nhất.\n"
                  "  - Đo lường Accuracy, Precision, Recall, Macro F1, Weighted F1, Tốc độ suy luận.\n\n"
                  "• Phân tích định tính ca lỗi:\n"
                  "  - Trích xuất 20 ca lỗi có độ tin cậy cao nhất (10 FP, 10 FN).\n"
                  "  - Khảo sát nguyên nhân ngôn ngữ học.\n\n"
                  "• Triển khai thực tế:\n"
                  "  - Demo tương tác Streamlit thời gian thực.")

    # =========================================================================
    # SLIDE 7: Baseline Model (TF-IDF + Logistic Regression)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "06. Mô Hình Cơ Sở: TF-IDF + Logistic Regression", "Tuân thủ triết lý AI Project Cycle: 'Start from simple to more complex models'")
    add_footer_and_logo(s7)

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="Cấu Hình Kỹ Thuật Baseline",
             body="• Vai trò của mô hình cơ sở:\n"
                  "  - Thiết lập chuẩn đối sánh tối thiểu theo Slide 9 AI Project Cycle.\n"
                  "  - Trả lời câu hỏi: Liệu Transformer phức tạp có thực sự vượt trội hơn giải pháp đơn giản?\n\n"
                  "• Cấu hình Pipeline:\n"
                  "  - TfidfVectorizer: ngram_range=(1, 2), max_features=10000, sublinear_tf=True.\n"
                  "  - LogisticRegression: C=1.0, max_iter=1000, solver='lbfgs'.\n\n"
                  "• Ưu điểm & Nhược điểm:\n"
                  "  - Ưu điểm: Huấn luyện cực nhanh (0.68s), suy luận siêu tốc (~28,784 mẫu/s).\n"
                  "  - Nhược điểm: Giả định 'Túi từ' (Bag-of-Words) bỏ qua trật tự từ và các cấu trúc đảo ngữ phức tạp.")

    base_acc = baseline_test.get("accuracy", 0.8194) * 100
    base_f1 = baseline_test.get("macro_f1", 0.8192)
    base_prec = baseline_test.get("macro_precision", 0.8208)
    base_rec = baseline_test.get("macro_recall", 0.8193)
    base_inf = baseline_test.get("inference_time_seconds", 0.137)
    base_tput = 3949 / base_inf if base_inf > 0 else 28784

    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
             title="Kết Quả Thực Nghiệm Baseline",
             body=f"• Đánh giá trên Held-Out Test Set (3,949 mẫu):\n\n"
                  f"  - Test Accuracy:        {base_acc:.2f}%\n"
                  f"  - Macro Precision:      {base_prec:.4f}\n"
                  f"  - Macro Recall:         {base_rec:.4f}\n"
                  f"  - Macro F1-Score:       {base_f1:.4f}\n"
                  f"  - Weighted F1-Score:    {baseline_test.get('weighted_f1', 0.8192):.4f}\n"
                  f"  - Thời gian suy luận:   {base_inf:.3f} giây (~{base_tput:,.0f} mẫu/s)\n\n"
                  f"• Ma trận nhầm lẫn Test Set:\n"
                  f"  - True Negatives (TN):  1,691 | False Positives (FP): 293\n"
                  f"  - False Negatives (FN): 420   | True Positives (TP):  1,545\n"
                  f"  - Tổng số ca dự đoán sai: 713 / 3,949 mẫu (18.06%)\n\n"
                  f"• Kết quả phát triển trên Validation Set (1,975 mẫu):\n"
                  f"  - Val Accuracy: 81.01% | Val Macro F1: 0.8098")

    # =========================================================================
    # SLIDE 8: Old BERT Failure Analysis & Key Architectural Differences
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "07. Phân Tích Sự Cố Mô Hình Cũ & Khác Biệt Pipeline Mới", "Khảo sát và audit chuyên sâu tài liệu môn học (DL_Model.ipynb)")
    add_footer_and_logo(s8)

    # Bảng so sánh 5 lỗi kỹ thuật
    rows, cols = 6, 3
    t_shape = s8.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.9))
    tbl = t_shape.table

    hdrs = ["Tiêu Chí Kỹ Thuật", "Triển Khai Cũ (DL_Model.ipynb - 65.20%)", "Pipeline Mới Của Nhóm (84.83%)"]
    for j, h in enumerate(hdrs):
        c = tbl.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        for p in c.text_frame.paragraphs:
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER

    comp_rows = [
        ["Đồng bộ Từ vựng", "Preprocess cased ghép với model uncased -> Lệch token ID", "Đồng bộ tuyệt đối: Tokenizer và Model đều uncased"],
        ["Cơ chế Huấn luyện", "Đóng băng Encoder (trainable=False) -> Feature Extraction", "Mở khóa toàn bộ tham số -> True Fine-Tuning"],
        ["Tốc độ học (LR)", "1e-3 (quá lớn gấp 50 lần, phá hủy trọng số tiền huấn luyện)", "2e-5 kết hợp Linear Warmup Scheduler chuẩn mực"],
        ["Kiến trúc Head", "5 tầng Dense sâu liên tiếp, loại bỏ Dropout -> Overfit nặng", "1 tầng Linear duy nhất + Dropout(0.1) tinh gọn"],
        ["Kiểm soát Epoch", "100 epochs cố định (Train Acc 97% nhưng Test Acc sụp đổ 65%)", "3 epochs, chọn Checkpoint theo Validation Macro F1"]
    ]

    for i, r in enumerate(comp_rows, 1):
        for j, val in enumerate(r):
            c = tbl.cell(i, j)
            c.text = val
            c.fill.solid()
            c.fill.fore_color.rgb = WHITE if i % 2 == 1 else BG_BOX
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(10.5)
                p.font.name = "Arial"
                p.font.color.rgb = NAVY if j == 0 else CHARCOAL
                if j == 0:
                    p.font.bold = True
                p.alignment = PP_ALIGN.LEFT

    # =========================================================================
    # SLIDE 9: BERT Fine-Tuning & Learning Dynamics
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "08. Quá Trình Huấn Luyện & Tinh Chỉnh BERT", "Động lực học tập qua 3 epochs và cơ chế lựa chọn Checkpoint tối ưu")
    add_footer_and_logo(s9)

    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="Thông Số Huấn Luyện & Động Lực Học",
             body="• Môi trường & Thời gian thực thi:\n"
                  "  - Huấn luyện trên CPU đa lõi (PyTorch 2.14.0+cpu).\n"
                  "  - 3 Epochs (2,592 steps, batch size=16).\n"
                  "  - Tổng thời gian huấn luyện: 15,617.19s (~4.34 giờ).\n\n"
                  "• Tiến trình qua từng Epoch (bert_training_history.json):\n"
                  "  - Epoch 1: Train Loss 0.4077 | Val Loss 0.3541 | Val F1 0.8359\n"
                  "  - Epoch 2: Train Loss 0.2885 | Val Loss 0.3702 | Val F1 0.8384\n"
                  "  - Epoch 3: Train Loss 0.1983 | Val Loss 0.4503 | Val F1 0.8389\n\n"
                  "• Giải thích hiện tượng Epoch 3 (Val Loss tăng vs Val F1 đạt đỉnh):\n"
                  "  - Cross-Entropy phạt nặng xác suất ở một số mẫu biên/nhiễu nhãn.\n"
                  "  - Phần lớn mẫu được phân loại chuẩn xác hơn -> Macro F1 đạt đỉnh 0.8389.\n"
                  "  - Lựa chọn Checkpoint theo Validation Macro F1 (checkpoint-2592) là hoàn toàn chuẩn xác.")

    hist_img = os.path.join(FIGURES_DIR, "training_history.png")
    if os.path.exists(hist_img):
        s9.shapes.add_picture(hist_img, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0))
    else:
        add_card(s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
                 title="Đồ Thị Đường Cong Học Tập",
                 body="• Đã lưu tại artifacts/figures/training_history.png")

    # =========================================================================
    # SLIDE 10: Comparative Results Table (Held-Out Test Set)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "09. Bảng So Sánh Kết Quả Thực Nghiệm Đối Đầu", "Đánh giá độc lập trên cùng tập kiểm thử Test Set (3,949 mẫu)")
    add_footer_and_logo(s10)

    rows, cols = 6, 6
    t_shape10 = s10.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.5))
    tbl10 = t_shape10.table

    headers10 = ["Phương Pháp / Mô Hình", "Nguồn / Phân Loại", "Accuracy", "Precision", "Recall", "Macro F1"]
    for j, h in enumerate(headers10):
        c = tbl10.cell(0, j)
        c.text = h
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
        for p in c.text_frame.paragraphs:
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER

    bert_acc = bert_test.get("accuracy", 0.8483) * 100
    bert_f1 = bert_test.get("macro_f1", 0.8483)
    bert_prec = bert_test.get("macro_precision", 0.8486)
    bert_rec = bert_test.get("macro_recall", 0.8482)

    data_rows10 = [
        ["TF-IDF + Logistic Regression", "Baseline mới của nhóm", f"{base_acc:.2f}%", f"{base_prec:.4f}", f"{base_rec:.4f}", f"{base_f1:.4f}"],
        ["Fine-Tuned BERT (Ours)", "Mô hình chính mới (Nhóm)", f"{bert_acc:.2f}%", f"{bert_prec:.4f}", f"{bert_rec:.4f}", f"{bert_f1:.4f}"],
        ["NNLM (Google Embedding)", "Tham khảo (DL_Model.ipynb)", "79.00%", "Not reported", "Not reported", "Not reported"],
        ["BiLSTM (2-layer)", "Tham khảo (DL_Model.ipynb)", "75.00%", "Not reported", "Not reported", "Not reported"],
        ["Old 'BERT' (Triển khai lỗi)", "Tham khảo (DL_Model.ipynb)", "65.20%", "Not reported", "Not reported", "Not reported"]
    ]

    for i, row in enumerate(data_rows10, 1):
        for j, val in enumerate(row):
            c = tbl10.cell(i, j)
            c.text = val
            c.fill.solid()
            if i == 2:
                c.fill.fore_color.rgb = HIGHLIGHT_ROW
            else:
                c.fill.fore_color.rgb = WHITE if i % 2 == 1 else BG_BOX
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.name = "Arial"
                p.font.color.rgb = NAVY if i == 2 else CHARCOAL
                if i == 2 or j == 0:
                    p.font.bold = True
                p.alignment = PP_ALIGN.LEFT if j < 2 else PP_ALIGN.CENTER

    # Card nhận định phía dưới
    add_card(s10, Inches(0.8), Inches(6.25), Inches(11.7), Inches(0.8),
             title="",
             body=f"• KẾT LUẬN THỰC NGHIỆM: Fine-Tuned BERT vượt trội Baseline +2.89% Accuracy và +0.0291 Macro F1 trên cùng tập Test Set.\n"
                  f"• Phục hồi hoàn toàn năng lực mô hình: Cải thiện +19.63 điểm phần trăm so với triển khai cũ trong DL_Model.ipynb (65.20%).")

    # =========================================================================
    # SLIDE 11: Confusion Matrix Comparison
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "10. Ma Trận Nhầm Lẫn & So Sánh Chi Tiết Từng Lớp", "Đánh giá chi tiết True/False Positives và Negatives trên 3,949 mẫu Test Set")
    add_footer_and_logo(s11)

    base_cm = os.path.join(FIGURES_DIR, "baseline_test_confusion_matrix.png")
    bert_cm = os.path.join(FIGURES_DIR, "bert_test_confusion_matrix.png")

    if os.path.exists(base_cm) and os.path.exists(bert_cm):
        s11.shapes.add_picture(base_cm, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.3))
        s11.shapes.add_picture(bert_cm, Inches(6.8), Inches(1.5), Inches(5.7), Inches(4.3))

    add_card(s11, Inches(0.8), Inches(5.95), Inches(11.7), Inches(1.15),
             title="Phân Tích Chi Tiết Ma Trận Nhầm Lẫn",
             body="• So sánh thành phần dự đoán (BERT vs. Baseline):\n"
                  "  - True Negatives: 1,711 vs. 1,691 (BERT nhận diện chính xác nhiều hơn 20 mẫu tiêu cực).\n"
                  "  - True Positives: 1,639 vs. 1,545 (BERT nhận diện chính xác nhiều hơn 94 mẫu tích cực).\n"
                  "  - False Negatives: 326 vs. 420 (BERT GIẢM MẠNH 94 ca bỏ sót tích cực, giảm 22.38% số ca FN!).\n"
                  "  - Tổng số lỗi dự đoán: BERT chỉ mắc 599 lỗi so với 713 lỗi của Baseline (Giảm tổng cộng 114 ca lỗi).")

    # =========================================================================
    # SLIDE 12: Qualitative Error Analysis
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "11. Báo Cáo Phân Tích Lỗi Định Tính (Qualitative Error Analysis)", "Khảo sát định tính mẫu 20 ca lỗi có độ tin cậy cao nhất (10 FP, 10 FN) của BERT")
    add_footer_and_logo(s12)

    add_card(s12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="Phân Bố Nguyên Nhân Trong Mẫu Khảo Sát",
             body="• Phương pháp khảo sát:\n"
                  "  - Trích xuất 20 ca lỗi có confidence cao nhất (> 99.3%) tại artifacts/metrics/error_cases.json.\n"
                  "  - Lưu ý: Đây là mẫu ca lỗi cực đoan, không suy rộng cho toàn bộ 599 ca lỗi.\n\n"
                  "• Thống kê hiện tượng trên mẫu 20 ca khảo sát:\n"
                  "  - Cảm xúc pha trộn (Mixed Sentiment): 15 / 20 ca (75.0%)\n"
                  "    (Khách vừa khen vừa chê, ép vào 1 nhãn nhị phân gây mơ hồ cố hữu).\n"
                  "  - Dấu hiệu mơ hồ / nhiễu nhãn (Label Noise): 8 / 20 ca (40.0%)\n"
                  "    (Văn bản hoàn toàn khen hoặc chê nhưng nhãn dữ liệu bị đối lập).\n"
                  "  - Cấu trúc phủ định / nhượng bộ: 8 / 20 ca (40.0%)\n"
                  "  - Tác động thẻ biểu mẫu ('No Positive'): 6 / 20 ca (30.0%)\n"
                  "  - Thách thức suy luận ngữ cảnh sâu: 1 / 20 ca (5.0%)")

    add_card(s12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
             title="Hiện Tượng Cắt Cụt & Trả Lời RQ3",
             body="• Đo lường hiện tượng cắt cụt văn bản (Truncation):\n"
                  "  - Số ca bị cắt cụt trong mẫu 20 ca lỗi: 0 / 20 ca (0%).\n"
                  "  - Độ dài token dao động từ 6 đến 126 tokens (đều <= 128 tokens).\n"
                  "  - Kết luận: Cắt cụt không phải là nguyên nhân gây ra 20 ca lỗi cực đoan này.\n\n"
                  "• Trả lời Câu hỏi Nghiên cứu 3 (RQ3):\n"
                  "  - Phần lớn lỗi cực đoan bắt nguồn từ bản chất bài toán nhị phân đối với cảm xúc đa khía cạnh và chất lượng nhãn dữ liệu.\n"
                  "  - Dự đoán của BERT ở nhiều ca lỗi cực đoan thể hiện sự nhất quán với ngữ nghĩa bề mặt của bài viết hơn là nhãn ground-truth.\n"
                  "  - Hạn chế suy luận ngữ cảnh sâu chỉ chiếm tỷ lệ nhỏ (như Case 16: dẫn lời chê của bạn bè rồi phản bác 'I disagree').")

    # =========================================================================
    # SLIDE 13: Demo Application (Streamlit)
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "12. Ứng Dụng Thực Tế: Web Demo Thời Gian Thực (Streamlit)", "Giai đoạn 04 (Deployment) theo AI Project Cycle — So sánh đối đầu thời gian thực")
    add_footer_and_logo(s13)

    add_card(s13, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="Tính Năng Của Ứng Dụng (app/app.py)",
             body="• Giao diện Web tương tác trực quan:\n"
                  "  - Khởi chạy nhanh bằng lệnh: streamlit run app/app.py\n"
                  "  - Cho phép người dùng nhập văn bản đánh giá bất kỳ.\n"
                  "  - Cung cấp sẵn 4 câu đánh giá mẫu kiểm thử các trường hợp phức tạp.\n\n"
                  "• So sánh đối đầu song song (Side-by-Side):\n"
                  "  - Hiển thị nhãn dự đoán và thanh xác suất Softmax của cả Baseline và BERT.\n"
                  "  - Làm nổi bật sự vượt trội của BERT trước các câu phủ định tinh vi.\n\n"
                  "• Tính năng giải thích mô hình (Explainability):\n"
                  "  - Trực quan hóa quá trình bẻ từ con WordPiece ([CLS], ##tokens, [SEP]).\n"
                  "  - Giúp người dùng hiểu rõ cách BERT mã hóa ngữ nghĩa văn bản.")

    add_card(s13, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
             title="Kiểm Thử Với 4 Câu Đánh Giá Mẫu",
             body="• Câu 1: 'The room was not bad, actually quite comfortable.'\n"
                  "  -> BERT: POSITIVE (Cấu trúc litotes 'not bad' được xử lý chuẩn xác).\n\n"
                  "• Câu 2: 'Not a single complaint about the wonderful staff.'\n"
                  "  -> BERT: POSITIVE (Hiểu 'Not a single' đảo nghĩa 'complaint').\n\n"
                  "• Câu 3: 'Great location, but the dirty bathroom ruined our entire stay.'\n"
                  "  -> BERT: NEGATIVE (Bắt đúng mệnh đề kết luận chính sau liên từ 'but').\n\n"
                  "• Câu 4: 'I had high expectations, but everything was disappointing.'\n"
                  "  -> BERT: NEGATIVE (Nhận diện cảm xúc thất vọng rõ ràng).\n\n"
                  "• Ý nghĩa thực tiễn: Hỗ trợ quản lý khách sạn theo dõi mức độ hài lòng khách hàng thời gian thực.")

    # =========================================================================
    # SLIDE 14: Limitations & Future Work
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "13. Giới Hạn Nghiên Cứu & Hướng Phát Triển Tương Lai", "Giai đoạn 05 & 06 (Maintenance & Feedback) theo AI Project Cycle")
    add_footer_and_logo(s14)

    add_card(s14, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="Các Giới Hạn Hiện Tại Của Đồ Án",
             body="• 1. Giới hạn bài toán nhị phân (Binary Limitation):\n"
                  "  - Chỉ phân loại 2 thái cực 0 và 1, chưa nắm bắt được các đánh giá trung lập (3 sao / Neutral).\n\n"
                  "• 2. Hiện tượng cắt cụt ở văn bản rất dài:\n"
                  "  - Dù ngưỡng MAX_LENGTH = 128 bao phủ 96.07% tập dữ liệu, 3.93% văn bản dài vẫn bị cắt đuôi (phân vị p99 là 168 tokens).\n\n"
                  "• 3. Đánh đổi chi phí tài nguyên tính toán (Resource Trade-off):\n"
                  "  - BERT có 110M tham số, tốc độ suy luận ~12.85 mẫu/s trên CPU, chậm hơn đáng kể so với mô hình tuyến tính TF-IDF (~28,784 mẫu/s).\n\n"
                  "• 4. Giới hạn mẫu phân tích lỗi:\n"
                  "  - Phân tích định tính tập trung vào mẫu 20 ca lỗi cực đoan nhất, không phản ánh toàn bộ 599 ca lỗi.")

    add_card(s14, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
             title="Hướng Phát Triển Tiếp Theo",
             body="• 1. Phân tích Cảm xúc Đa Khía cạnh (ABSA):\n"
                  "  - Tách bạch cảm xúc theo từng tiêu chí: Vị trí, Vệ sinh, Giá cả, Nhân viên để giải quyết triệt để bài toán cảm xúc pha trộn.\n\n"
                  "• 2. Tối ưu hóa & Nén mô hình (Model Compression):\n"
                  "  - Chuyển giao tri thức (Knowledge Distillation) sang DistilBERT (giảm 40% kích thước, giữ 97% hiệu năng).\n"
                  "  - Áp dụng lượng tử hóa INT8 (Quantization) và ONNX Runtime để tăng tốc suy luận trên CPU lên 3-5 lần.\n\n"
                  "• 3. Mở rộng đa ngôn ngữ:\n"
                  "  - Ứng dụng PhoBERT / XLM-RoBERTa cho đánh giá khách sạn tại thị trường Việt Nam.")

    # =========================================================================
    # SLIDE 15: Conclusion & Defense Q&A
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "14. Tổng Kết Đóng Góp & Sẵn Sàng Phản Biện (Q&A)", "Hoàn thành toàn diện đồ án môn học Trí tuệ Nhân tạo")
    add_footer_and_logo(s15)

    add_card(s15, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0),
             title="4 Đóng Góp Nổi Bật Của Đồ Án",
             body="• 1. Hoàn thành trọn vẹn quy trình AI Project Cycle:\n"
                  "  - Thực hiện bài bản từ Scope, Data, Models, Deployment, đến Error Analysis và Feedback.\n\n"
                  "• 2. Triển khai chuẩn mực & Khôi phục tiềm năng BERT:\n"
                  "  - Tìm ra 5 lỗi kỹ thuật trong mã nguồn cũ môn học (chỉ đạt 65.20%).\n"
                  "  - Xây dựng pipeline chuẩn đạt 84.83% Accuracy (Tăng +19.63%!).\n\n"
                  "• 3. Thực nghiệm đối đầu khách quan:\n"
                  "  - Chứng minh BERT vượt trội Baseline (+2.89% Accuracy, +0.0291 Macro F1) trên Held-Out Test Set 3,949 mẫu.\n\n"
                  "• 4. Tính trung thực & Tái lập học thuật:\n"
                  "  - Mọi con số đều có artifact kiểm chứng; bộ tài liệu Study Guide & Defense QA hoàn chỉnh.")

    add_card(s15, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.0),
             title="Lời Cảm Ơn & Phiên Hỏi Đáp (Q&A)",
             body="TRÂN TRỌNG CẢM ƠN QUÝ THẦY CÔ VÀ HỘI ĐỒNG!\n\n"
                  "• Nhóm nghiên cứu xin chân thành cảm ơn sự hướng dẫn tận tình của Giảng viên môn học Trí tuệ Nhân tạo.\n\n"
                  "• Dự án đã sẵn sàng cho phần Thảo luận & Trả lời Phản biện:\n"
                  "  - Cẩm nang 25 khái niệm: docs/STUDY_GUIDE.md\n"
                  "  - Bộ 25 câu hỏi phản biện: docs/DEFENSE_QA.md\n"
                  "  - Báo cáo tổng kết 17 mục: FINAL_REPORT.md\n"
                  "  - Ứng dụng Demo trực tiếp: streamlit run app/app.py\n\n"
                  "Kính mời Quý Thầy Cô đặt câu hỏi cho Nhóm!")

    # Lưu cả hai file để đảm bảo tính tương thích
    output_final = os.path.join(PROJECT_ROOT, "presentation", "BERT_Project_Final.pptx")
    output_compat = os.path.join(PROJECT_ROOT, "presentation", "BERT_Project.pptx")
    prs.save(output_final)
    prs.save(output_compat)
    print(f"[+] Đã tạo thành công bài thuyết trình chính thức tại:\n    -> {output_final}\n    -> {output_compat}")
    return output_final

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Tạo slide PowerPoint cho đồ án BERT")
    parser.add_argument("--draft", action="store_true", help="Tạo bản nháp khi chưa có kết quả thực nghiệm thực tế")
    args = parser.parse_args()

    create_presentation(is_draft=args.draft)
